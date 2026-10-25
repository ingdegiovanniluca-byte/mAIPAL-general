import asyncio, os, sys, copy
from datetime import datetime, timezone, timedelta
os.environ.update(MONGO_URL="mongodb://x", DB_NAME="t", OPENAI_API_KEY="x")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
import server, scheduled_actions as sa, list_updates as lu, usage_tracking as ut, retrieval

def _get(doc, k):
    return doc.get(k)

def _match(doc, q):
    for k, v in (q or {}).items():
        if k == "$or":
            if not any(_match(doc, s) for s in v): return False
            continue
        if k == "$and":
            if not all(_match(doc, s) for s in v): return False
            continue
        val = _get(doc, k)
        if isinstance(v, dict):
            for op, arg in v.items():
                if op == "$in" and not (any(x in arg for x in val) if isinstance(val, list) else val in arg): return False
                if op == "$nin" and val in arg: return False
                if op == "$ne" and val == arg: return False
                if op == "$gte" and (val is None or val < arg): return False
                if op == "$gt" and (val is None or val <= arg): return False
                if op == "$lte" and (val is None or val > arg): return False
                if op == "$lt" and (val is None or val >= arg): return False
                if op == "$exists" and ((k in doc) != arg): return False
        elif val != v:
            return False
    return True

class Cur:
    def __init__(s, d): s.d = d
    def sort(s, f, dr=1): s.d = sorted(s.d, key=lambda x: x.get(f) or "", reverse=dr == -1); return s
    async def to_list(s, n): return [copy.deepcopy(x) for x in s.d[:n]]
    def __aiter__(s):
        s._it = iter(copy.deepcopy(s.d)); return s
    async def __anext__(s):
        try: return next(s._it)
        except StopIteration: raise StopAsyncIteration

class R:
    def __init__(s, n): s.modified_count = n; s.deleted_count = n

class Coll:
    def __init__(s): s.docs = []
    def find(s, q=None, proj=None): return Cur([d for d in s.docs if _match(d, q)])
    async def find_one(s, q=None, proj=None):
        for d in s.docs:
            if _match(d, q): return copy.deepcopy(d)
    async def insert_one(s, d): s.docs.append(copy.deepcopy(d))
    async def count_documents(s, q): return sum(1 for d in s.docs if _match(d, q))
    async def update_one(s, q, u, upsert=False):
        for d in s.docs:
            if _match(d, q):
                d.update(u.get("$set", {}))
                for k in u.get("$unset", {}): d.pop(k, None)
                for k, v in u.get("$addToSet", {}).items():
                    arr = d.setdefault(k, [])
                    if v not in arr: arr.append(v)
                for k, v in u.get("$pull", {}).items():
                    out = v.get("$in") if isinstance(v, dict) else [v]
                    d[k] = [x for x in d.get(k, []) if x not in out]
                for k, v in u.get("$inc", {}).items(): d[k] = (d.get(k) or 0) + v
                for k, v in u.get("$push", {}).items():
                    arr = d.setdefault(k, [])
                    if isinstance(v, dict) and "$each" in v:
                        arr.extend(v["$each"]); 
                        if "$slice" in v: d[k] = arr[v["$slice"]:]
                    else: arr.append(v)
                return R(1)
        if upsert:
            d = {k: v for k, v in q.items() if not isinstance(v, dict)}; d.update(u.get("$set", {})); s.docs.append(d)
        return R(0)
    async def insert_many(s, ds): s.docs.extend(copy.deepcopy(d) for d in ds)
    async def update_many(s, q, u):
        n = 0
        for d in s.docs:
            if _match(d, q):
                d.update(u.get("$set", {})); n += 1
                for k, v in u.get("$pull", {}).items():
                    out = v.get("$in") if isinstance(v, dict) else [v]
                    d[k] = [x for x in d.get(k, []) if x not in out]
        return R(n)
    async def delete_one(s, q):
        for i, d in enumerate(s.docs):
            if _match(d, q): s.docs.pop(i); return R(1)
        return R(0)
    async def delete_many(s, q):
        n = len(s.docs); s.docs = [d for d in s.docs if not _match(d, q)]; return R(n - len(s.docs))

class DB:
    def __getattr__(s, n):
        c = Coll(); setattr(s, n, c); return c

db = DB()
server.db = db
ut.fire_and_forget_feature_event = lambda **k: None
ut.fire_and_forget_llm_call = lambda **k: None
from llm_integrations import TextDelta, StreamDone
import job_logs as jl, telegram_bot as tg, retrieval, embeddings as emb
import scheduled_actions as sa
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
import io, work_reports as wr
from datetime import date
from docx import Document
from fastapi import HTTPException

U = {"user_id": "u1", "email": "a@b.it", "name": "Luca Rossi", "role": "admin"}
db.users.docs.append(dict(U))
SAVED = []
async def fake_deliver(current, docx_bytes, fname, folder, subs, caption):
    SAVED.append({"folder": folder, "subs": subs, "fname": fname, "bytes": docx_bytes}); return "https://drive/x", "Bagno", False
server._deliver_work_report = fake_deliver
ASKED = []
async def fake_summary(commessa, period, diary, facts, wants, request="", user_id=None, channel="web"):
    ASKED.append({"period": period, "diary": diary, "facts": facts, "wants": wants, "request": request})
    out = {"stato": "Impianti completati, in attesa delle piastrelle.", "attivita": "- Posa tubazioni\n- Tracce"}
    return {k: out[k] for k in ("stato", "attivita") if k in wants}
wr.commessa_summary = fake_summary

def doc_text(b):
    d = Document(io.BytesIO(b))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts += [c.text for c in r.cells]
    return "\n".join(parts)

def log(cid, day, hours=(), mats=(), problems=(), text="lavori", kind=None):
    db.job_logs.docs.append({"id": f"jl_{len(db.job_logs.docs)}", "commessa_id": cid, "date": day, "kind": kind, "author_name": "Luca",
                             "text": text, "raw_text": text, "hours": [{"who": w, "hours": h} for w, h in hours],
                             "materials": [{"name": n, "qty": q, "unit": u} for n, q, u in mats], "problems": list(problems),
                             "photos": [], "created_at": day + "T10:00:00+00:00"})

async def main():
    T = date(2026, 10, 24)
    # 1) the period said in the request, against a fixed "today"
    P = lambda s: wr.parse_period(s, T)
    assert P("report della commessa Rossi") is None
    assert P("dal 1/10 al 15/10") == (date(2026, 10, 1), date(2026, 10, 15))
    assert P("dal 3 al 20 settembre") == (date(2026, 9, 3), date(2026, 9, 20))
    assert P("dal 5 settembre") == (date(2026, 9, 5), T)
    assert P("dal 20/12") == (date(2025, 12, 20), T)                                   # a future day: last year's
    assert P("report di ottobre") == (date(2026, 10, 1), T)                           # up to today, not the month's end
    assert P("di settembre") == (date(2026, 9, 1), date(2026, 9, 30))
    assert P("novembre") == (date(2025, 11, 1), date(2025, 11, 30))                   # a month still to come: last year's
    assert P("il 3 ottobre abbiamo finito") is None                                     # a day in a sentence is not a month
    assert P("ultimi 15 giorni") == (date(2026, 10, 10), T) and P("ultime due settimane") == (date(2026, 10, 11), T)
    assert P("ultima settimana") == (date(2026, 10, 18), T) and P("ultimo mese") == (date(2026, 9, 25), T)
    assert P("mese scorso") == (date(2026, 9, 1), date(2026, 9, 30)) and P("questo mese") == (date(2026, 10, 1), T)
    assert P("settimana scorsa") == (date(2026, 10, 12), date(2026, 10, 18))           # Monday to Sunday
    assert P("ieri") == (date(2026, 10, 23),) * 2
    assert wr.parts_from_text("solo ore e problemi") == ["ore", "problemi"]
    assert wr.parts_from_text("costruzione del muro") == [] and wr.parts_from_text("i costi e lo stato della commessa") == ["stato", "costi"]
    assert wr.period_label(date(2026, 10, 1), date(2026, 10, 15)) == "dal 01/10/2026 al 15/10/2026"

    me = await server.update_profile(server.ProfilePatch(business_vertical="artigiano"), server.User(**U))
    cl = next(c for c in db.collections.docs if c.get("system_key") == "artigiano_clienti")
    db.collection_items.docs.append({"id": "c_rossi", "collection_id": cl["id"], "data": {"nome": "Mario Rossi", "indirizzo": "Via Roma 5"}})
    db.collection_items.docs.append({"id": "c_bianchi", "collection_id": cl["id"], "data": {"nome": "Anna Bianchi"}})
    for i, it, t, st, start in (("s_bagno", "c_rossi", "Bagno", "in corso", "2026-09-01"), ("s_cucina", "c_rossi", "Cucina", "in corso", ""),
                                ("s_tetto", "c_bianchi", "Tetto", "in corso", "")):
        db.collection_sub_items.docs.append({"id": i, "collection_id": cl["id"], "item_id": it, "data": {"titolo": t, "stato": st, "data_inizio": start}})
    log("s_bagno", "2026-09-10", hours=[("Luca", 8)], mats=[("corrugato 25", 20, "m")], problems=["tubo rotto sotto il lavello"], text="tracce")
    log("s_bagno", "2026-10-05", hours=[("Luca", 6), ("Gino", 4)], problems=["piastrelle in ritardo"], text="posa tubazioni")
    log("s_bagno", "2026-10-06", text="Stato: sospesa", kind="stato")
    log("s_cucina", "2026-10-05", hours=[("Luca", 3)], text="cucina")
    db.job_milestones.docs.append({"id": "ms1", "commessa_id": "s_bagno", "title": "Impianti", "due_date": "2026-10-01",
                                   "done": True, "done_date": "2026-10-02", "created_at": "x"})
    jl.today_local = lambda: T
    await server.patch_job_costs(server.CostsPatch(default_rate=30), "s_bagno", me)
    db.job_economics.docs.append({"commessa_id": "s_bagno", "budget": 5000, "alerts": [],
                                  "extras": [{"id": "x1", "label": "Ponteggio", "amount": 200, "date": "2026-10-03"},
                                             {"id": "x2", "label": "Smaltimento", "amount": 50, "date": "2026-09-02"}]})

    # 2) the templates list it, with its parts; the five dictated models stay as they are
    t = await server.list_work_templates(me)
    assert [x["key"] for x in t["builtin"]][-1] == "commessa" and t["builtin"][-1]["commessa"]
    assert [p["key"] for p in t["builtin"][-1]["parts"]] == ["stato", "attivita", "ore", "problemi", "materiali", "milestone", "costi"]
    assert list(wr.BUILTIN_TEMPLATES) == ["sopralluogo", "intervento", "sal", "fine_lavori", "variante"]

    # 3) a client with two commesse: asked which; nothing made, no model call
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(text="report dei Rossi"))
    assert r["status"] == "pick_commessa" and {c["commessa_id"] for c in r["candidates"]} == {"s_bagno", "s_cucina"} and not ASKED
    nsubs, nitems = len(db.collection_sub_items.docs), len(db.collection_items.docs)

    # 4) picked, October, hours + problems + costs + state: only October's entries
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(
        text="report di ottobre, evidenzia i ritardi", commessa_id="s_bagno", include=["costi", "ore", "problemi", "stato", "boh"]))
    assert r["status"] == "ok" and r["report_type"] == "commessa" and r["commessa_id"] == "s_bagno" and r["client_name"] == "Mario Rossi"
    assert r["period_label"] == "dal 01/10/2026 al 24/10/2026" and r["include"] == ["costi", "ore", "problemi", "stato"]
    assert [s["title"] for s in r["sections"]] == ["Dati della commessa", "Riassunto stato commessa", "Ore lavorate", "Problematiche", "Costi"]
    sec = {s["title"]: s["fields"] for s in r["sections"]}
    assert sec["Dati della commessa"]["Periodo di riferimento"] == "dal 01/10/2026 al 24/10/2026" and sec["Dati della commessa"]["Commessa"] == "Bagno"
    assert sec["Ore lavorate"]["Ore totali"] == "10 h" and sec["Ore lavorate"]["Ore per persona"] == "- Luca: 6 h\n- Gino: 4 h"
    assert sec["Problematiche"]["Problemi riscontrati"] == "- 05/10/2026: piastrelle in ritardo"            # September's stays out
    assert sec["Costi"]["Manodopera"] == "€ 300" and sec["Costi"]["Altre spese"] == "€ 200" and sec["Costi"]["Totale costi"] == "€ 500"
    assert "Preventivo" not in sec["Costi"]                                                         # a period: no margin on the whole quote
    assert sec["Riassunto stato commessa"]["Riassunto"] == "Impianti completati, in attesa delle piastrelle."
    a = ASKED[-1]
    assert a["wants"] == ["costi", "ore", "problemi", "stato"] and "evidenzia i ritardi" in a["request"]
    assert "posa tubazioni" in a["diary"] and "tracce" not in a["diary"] and "Stato: sospesa" in a["diary"] and "Impianti" in a["facts"]
    txt = doc_text(SAVED[-1]["bytes"])
    for must in ("Report di commessa", "Mario Rossi", "Bagno", "piastrelle in ritardo", "Firma del tecnico"):
        assert must in txt, must
    assert "tubo rotto" not in txt and SAVED[-1]["folder"] == "Clienti" and SAVED[-1]["subs"] == ["Mario Rossi", "Bagno"]
    assert len(db.collection_sub_items.docs) == nsubs and len(db.collection_items.docs) == nitems    # nothing created
    line = [x for x in db.job_logs.docs if x.get("kind") == "report"][-1]
    assert line["commessa_id"] == "s_bagno" and line["text"] == "Report «Report di commessa» (dal 01/10/2026 al 24/10/2026) generato da Luca."
    stored = db.work_reports.docs[-1]
    assert stored["docx_b64"] and stored["period_from"] == "2026-10-01" and stored["summary"].startswith("Impianti completati")

    # 5) the commessa named, the whole job (from its start date), the parts said in the text; no model call
    n = len(ASKED)
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(text="Bagno dei Rossi: materiali, milestone e costi"))
    assert r["commessa_id"] == "s_bagno" and r["include"] == ["materiali", "milestone", "costi"] and len(ASKED) == n
    assert r["period_label"] == "dal 01/09/2026 al 24/10/2026"
    sec = {s["title"]: s["fields"] for s in r["sections"]}
    assert sec["Materiali"]["Materiali impiegati"] == "- 20 m corrugato 25" and sec["Milestone"]["Milestone"] == "- Impianti: completata il 02/10/2026"
    assert sec["Costi"]["Manodopera"] == "€ 540" and sec["Costi"]["Totale costi"] == "€ 790" and sec["Costi"]["Preventivo"] == "€ 5.000"
    assert sec["Costi"]["Non conteggiati"] == "1 materiali senza prezzo"

    # 6) the chips' periods; nothing chosen and nothing said -> state, hours, problems; an empty period says so
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(text="", commessa_id="s_bagno", period="prev_month"))
    assert r["period_label"] == "dal 01/09/2026 al 30/09/2026" and r["include"] == ["stato", "ore", "problemi"]
    assert {s["title"]: s["fields"] for s in r["sections"]}["Ore lavorate"]["Ore totali"] == "8 h"
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno", period="custom",
                                                                                date_from="2026-10-20", date_to="2026-10-10", include=["ore", "problemi"]))
    assert r["period_label"] == "dal 10/10/2026 al 20/10/2026"
    sec = {s["title"]: s["fields"] for s in r["sections"]}
    assert sec["Ore lavorate"]["Ore totali"] == "0 h" and sec["Problematiche"]["Problemi riscontrati"] == "Nessun problema registrato nel periodo"
    assert "Ore lavorate" in doc_text(SAVED[-1]["bytes"])                                           # a chosen part is in the file even when empty
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno", period="custom", date_to="2026-09-30", include=["ore"]))
    assert r["period_label"] == "dal 01/09/2026 al 30/09/2026"                                       # only "up to": from the start
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno", period="last7", include=["ore"]))
    assert r["period_label"] == "dal 18/10/2026 al 24/10/2026"
    try:
        await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno", period="custom", date_from="x"))
        assert False
    except HTTPException as ex:
        assert ex.status_code == 400

    # 7) the model fails: the work done comes from the diary's own lines, the report is still made
    async def broken(*a, **k): raise RuntimeError("down")
    wr.commessa_summary = broken
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno", include=["attivita"], period="last30"))
    assert {s["title"]: s["fields"] for s in r["sections"]}["Lavori svolti"]["Lavori del periodo"] == "- 05/10/2026 · Luca: posa tubazioni"
    wr.commessa_summary = fake_summary

    # 8) a report di commessa is changed like the others, in the same chat
    async def fake_revise(sections, instruction, user_id=None, channel="web"):
        out = [{"title": s["title"], "fields": dict(s["fields"])} for s in sections]
        out[1]["fields"]["Riassunto"] = "Lavori sospesi."
        return {"sections": out, "client": {"name": "", "phone": "", "email": "", "address": ""}, "job_title": ""}
    wr.revise = fake_revise
    r = await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_bagno"))
    rv = await server._revise_work_report(me, r["id"], "scrivi che i lavori sono sospesi")
    assert rv["changed"] == ["Riassunto stato commessa › Riassunto"] and rv["version"] == 2

    # 8b) the request written with "Primo sopralluogo" still selected: a report di commessa, never an empty sopralluogo
    assert wr.asks_commessa_report("Fammi il report della commessa impianto elettrico") and wr.asks_commessa_report("report di cantiere")
    assert not wr.asks_commessa_report("sopralluogo dai Rossi per la commessa del bagno")
    assert P("per la settimana corrente") == (date(2026, 10, 19), T) and P("mese corrente") == (date(2026, 10, 1), T)
    db.collection_items.docs.append({"id": "c_lucad", "collection_id": cl["id"], "data": {"nome": "Luca D"}})
    db.collection_sub_items.docs.append({"id": "s_imp", "collection_id": cl["id"], "item_id": "c_lucad", "data": {"titolo": "Impianto elettrico", "stato": "in corso"}})
    log("s_imp", "2026-10-20", hours=[("Luca", 8), ("Gino", 6)], mats=[("cavo 2,5", 50, "m")], text="passaggio cavi")
    log("s_imp", "2026-10-12", hours=[("Luca", 5)], text="tracce (settimana prima)")
    async def no_dictation(*a, **k): raise AssertionError("not a dictated report")
    wr.interpret = no_dictation
    nsubs = len(db.collection_sub_items.docs)
    msg = ("Fammi il report della commessa impianto elettrico del cliente luca d per la settimana corrente. "
           "includi attività fatte, materiali usati e ora lavoarate divise per lavoratore.")
    r = await server._generate_work_report(me, msg)   # the type selected in the chat: Primo sopralluogo
    assert r["report_type"] == "commessa" and r["commessa_id"] == "s_imp" and r["client_name"] == "Luca D"
    assert r["include"] == ["attivita", "ore", "materiali"] and r["period_label"] == "dal 19/10/2026 al 24/10/2026"
    sec = {s["title"]: s["fields"] for s in r["sections"]}
    assert sec["Ore lavorate"]["Ore per persona"] == "- Luca: 8 h\n- Gino: 6 h" and sec["Materiali"]["Materiali impiegati"] == "- 50 m cavo 2,5"
    assert "Lavori svolti" in sec and len(db.collection_sub_items.docs) == nsubs

    assert r["type_from_text"]   # Primo sopralluogo was selected, the message asked for this one

    # 8c) the message decides the type, the selection only when it says none
    D = wr.detect_report_type
    assert D("SAL del bagno dei Rossi: fatto il 60%") == "sal" and D("stato di avanzamento lavori Bianchi") == "sal"
    assert D("Fine lavori del bagno dei Rossi, collaudato") == "fine_lavori" and D("verbale di consegna ai Verdi") == "fine_lavori"
    assert D("Variante per i Rossi: una presa in più") == "variante" and D("lavori extra dai Bianchi") == "variante"
    assert D("Intervento dai Verdi: riparato il quadro") == "intervento" and D("riparazione caldaia Neri") == "intervento"
    assert D("sopralluogo dai Rossi per una variante in bagno") == "sopralluogo" and D("sopraluogo Bianchi") == "sopralluogo"
    assert D("report della commessa bagno") == "commessa" and D("dai Rossi, bagno da rifare") is None
    assert D("salvami il bagno dei Rossi") is None and D("SAL.") == "sal"
    assert D("compila il Verbale ASL per i Rossi", [{"id": "wtpl_1", "name": "Verbale ASL"}]) == "wtpl_1"
    # the chips with "Report di commessa" selected; what the message says wins over them
    r = await server._generate_work_report(me, "report della commessa Bagno dei Rossi: solo i materiali, mese scorso", "commessa",
                                           commessa_opts={"include": ["ore"], "period": "last7", "date_from": None, "date_to": None})
    assert r["report_type"] == "commessa" and r["include"] == ["materiali"] and r["period_label"] == "dal 01/09/2026 al 30/09/2026"
    assert not r["type_from_text"]
    r = await server._generate_work_report(me, "Bagno dei Rossi", "commessa",
                                           commessa_opts={"include": ["ore"], "period": "last7", "date_from": None, "date_to": None})
    assert r["include"] == ["ore"] and r["period_label"] == "dal 18/10/2026 al 24/10/2026"   # nothing said: the chips

    # 9) a client with no commessa / an unknown commessa
    db.collection_items.docs.append({"id": "c_verdi", "collection_id": cl["id"], "data": {"nome": "Paolo Verdi"}})
    try:
        await server._generate_commessa_report(me, server.CommessaReportPayload(text="report di Verdi")); assert False
    except HTTPException as ex:
        assert ex.status_code == 400 and "Paolo Verdi non ha ancora commesse" in ex.detail
    try:
        await server._generate_commessa_report(me, server.CommessaReportPayload(commessa_id="s_nope")); assert False
    except HTTPException as ex:
        assert ex.status_code == 404
    assert not [s for s in db.collection_sub_items.docs if s["item_id"] == "c_verdi"]
    print("ALL OK")
asyncio.run(main())
