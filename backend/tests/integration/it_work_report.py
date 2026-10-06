import asyncio, os, sys, copy
from datetime import datetime, timezone, timedelta
os.environ.update(MONGO_URL="mongodb://x", DB_NAME="t", OPENAI_API_KEY="x")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
import server, scheduled_actions as sa, list_updates as lu, usage_tracking as ut, retrieval

def _get(doc, k):
    return doc.get(k)

def _match(doc, q):
    for k, v in (q or {}).items():
        if k == "$or":
            if not any(_match(doc, s) for s in v): return False
            continue
        if k == "$and":
            if not all(_match(doc, s) for s in v): return False
            continue
        val = _get(doc, k)
        if isinstance(v, dict):
            for op, arg in v.items():
                if op == "$in" and val not in arg: return False
                if op == "$nin" and val in arg: return False
                if op == "$ne" and val == arg: return False
                if op == "$gte" and (val is None or val < arg): return False
                if op == "$gt" and (val is None or val <= arg): return False
                if op == "$lte" and (val is None or val > arg): return False
                if op == "$lt" and (val is None or val >= arg): return False
                if op == "$exists" and ((k in doc) != arg): return False
        elif val != v:
            return False
    return True

class Cur:
    def __init__(s, d): s.d = d
    def sort(s, f, dr=1): s.d = sorted(s.d, key=lambda x: x.get(f) or "", reverse=dr == -1); return s
    async def to_list(s, n): return [copy.deepcopy(x) for x in s.d[:n]]
    def __aiter__(s):
        s._it = iter(copy.deepcopy(s.d)); return s
    async def __anext__(s):
        try: return next(s._it)
        except StopIteration: raise StopAsyncIteration

class R:
    def __init__(s, n): s.modified_count = n; s.deleted_count = n; s.matched_count = n

class Coll:
    def __init__(s): s.docs = []
    def find(s, q=None, proj=None): return Cur([d for d in s.docs if _match(d, q)])
    async def find_one(s, q=None, proj=None):
        for d in s.docs:
            if _match(d, q): return copy.deepcopy(d)
    async def insert_one(s, d): s.docs.append(copy.deepcopy(d))
    async def count_documents(s, q): return sum(1 for d in s.docs if _match(d, q))
    async def update_one(s, q, u, upsert=False):
        for d in s.docs:
            if _match(d, q):
                d.update(u.get("$set", {}))
                for k, v in u.get("$inc", {}).items(): d[k] = (d.get(k) or 0) + v
                for k, v in u.get("$addToSet", {}).items():
                    arr = d.setdefault(k, [])
                    for x in (v["$each"] if isinstance(v, dict) and "$each" in v else [v]):
                        if x not in arr: arr.append(x)
                for k, v in u.get("$push", {}).items():
                    arr = d.setdefault(k, [])
                    if isinstance(v, dict) and "$each" in v:
                        arr.extend(v["$each"]); 
                        if "$slice" in v: d[k] = arr[v["$slice"]:]
                    else: arr.append(v)
                return R(1)
        if upsert:
            d = {k: v for k, v in q.items() if not isinstance(v, dict)}; d.update(u.get("$set", {})); s.docs.append(d)
        return R(0)
    async def insert_many(s, ds): s.docs.extend(copy.deepcopy(d) for d in ds)
    async def update_many(s, q, u):
        n = 0
        for d in s.docs:
            if _match(d, q): d.update(u.get("$set", {})); n += 1
        return R(n)
    async def delete_one(s, q):
        for i, d in enumerate(s.docs):
            if _match(d, q): s.docs.pop(i); return R(1)
        return R(0)
    async def delete_many(s, q):
        n = len(s.docs); s.docs = [d for d in s.docs if not _match(d, q)]; return R(n - len(s.docs))

class DB:
    def __getattr__(s, n):
        c = Coll(); setattr(s, n, c); return c

db = DB()
server.db = db
ut.fire_and_forget_feature_event = lambda **k: None
ut.fire_and_forget_llm_call = lambda **k: None
from llm_integrations import TextDelta, StreamDone
import json as J
U = {"user_id": "u1", "email": "a@b.it", "name": "Luca"}
db.users.docs.append(U)




from fastapi import HTTPException
async def no_sharing(colls, current): return colls
server._add_sharing_info = no_sharing
class Agg:
    def __init__(s, out): s.out = out
    async def to_list(s, n): return s.out
db.collection_items.aggregate = lambda pipe: Agg([])

import io, work_reports as wr
from docx import Document

SAVED = []
async def fake_targets(uid): return ["google"]
async def fake_save(uid, folder, tmp, fname, ctype, only=None, subfolders=None):
    SAVED.append({"folder": folder, "subs": subfolders, "fname": fname, "bytes": open(tmp, "rb").read()})
    return {"saved": ["google"], "web_view_link": "https://drive/x", "links": {}}
server._storage_targets = fake_targets
server._storage_save = fake_save

CALLS = []
def interp(client, job, fields=None):
    async def f(text, sections, user_name="", user_id=None, channel="web"):
        CALLS.append(text)
        out = [{"title": s["title"], "fields": {k: (fields or {}).get(k, "") for k in s["fields"]}} for s in sections]
        return {"sections": out, "client": {"name": "", "phone": "", "email": "", "address": "", **client}, "job_title": job}
    return f

def doc_text(b):
    d = Document(io.BytesIO(b))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts += [c.text for c in r.cells]
    return "\n".join(parts)

async def main():
    u = server.User(**U)
    me = await server.update_profile(server.ProfilePatch(business_vertical="artigiano"), u)
    cl = next(c for c in db.collections.docs if c.get("system_key") == "artigiano_clienti")

    # 1) new client + new commessa from the dictation
    wr.interpret = interp({"name": "Mario Rossi", "phone": "333 1234567", "address": "Via Roma 5, Torino"},
                          "Rifacimento impianto elettrico cucina",
                          {"Oggetto dell'intervento": "Rifacimento impianto elettrico della cucina",
                           "Descrizione dello stato attuale": "Impianto non a norma, senza differenziale.",
                           "Lavori proposti": "- Nuovo quadro\n- 6 punti presa", "Cliente": "Mario Rossi"})
    r = await server._generate_work_report(me, "sopralluogo dal signor Mario Rossi in via Roma 5 per l'impianto della cucina")
    print({k: r[k] for k in ("client_name", "client_created", "commessa_title", "commessa_created", "drive_folder")})
    assert r["status"] == "ok" and r["client_created"] and r["commessa_created"] and "docx_b64" not in r
    items = [i for i in db.collection_items.docs if i["collection_id"] == cl["id"]]
    assert len(items) == 1 and items[0]["data"] == {"nome": "Mario Rossi", "telefono": "333 1234567", "indirizzo": "Via Roma 5, Torino"}
    subs = db.collection_sub_items.docs
    assert len(subs) == 1 and subs[0]["item_id"] == items[0]["id"]
    assert subs[0]["data"]["titolo"] == "Rifacimento impianto elettrico cucina" and subs[0]["data"]["stato"] == "preventivo"
    assert subs[0]["data"]["indirizzo_cantiere"] == "Via Roma 5, Torino"
    assert SAVED[-1]["folder"] == "Clienti" and SAVED[-1]["subs"] == ["Mario Rossi", "Rifacimento impianto elettrico cucina"]
    txt = doc_text(SAVED[-1]["bytes"])
    for must in ("Primo sopralluogo", "Mario Rossi", "333 1234567", "Via Roma 5, Torino", "Luca", "Firma del tecnico",
                 "Firma del cliente per presa visione", "Luogo e data", "Impresa di Luca", "Nuovo quadro"):
        assert must in txt, must
    assert "Note e prescrizioni" not in txt   # a section nobody talked about stays out
    stored = db.work_reports.docs[-1]
    assert stored["docx_b64"] and stored["client_item_id"] == items[0]["id"] and stored["commessa_id"] == subs[0]["id"]

    # 2) same client named again, same commessa title in the text: nothing duplicated; missing email filled in
    wr.interpret = interp({"name": "Rossi", "email": "mario@rossi.it"}, "Impianto cucina")
    r = await server._generate_work_report(me, "Torno da Rossi per il rifacimento impianto elettrico cucina, email mario@rossi.it")
    assert not r["client_created"] and not r["commessa_created"] and r["commessa_id"] == subs[0]["id"]
    items = [i for i in db.collection_items.docs if i["collection_id"] == cl["id"]]
    assert len(items) == 1 and items[0]["data"]["email"] == "mario@rossi.it" and items[0]["data"]["telefono"] == "333 1234567"

    # 3) a second Rossi -> ambiguous, no model call; then picked / new client
    db.collection_items.docs.append({"id": "item_r2", "collection_id": cl["id"], "data": {"nome": "Anna Rossi", "indirizzo": "Corso Francia 1"}})
    n = len(CALLS)
    r = await server._generate_work_report(me, "sopralluogo da Rossi per il bagno")
    print(r["status"], [c["name"] for c in r["candidates"]])
    assert r["status"] == "ambiguous_client" and len(r["candidates"]) == 2 and len(CALLS) == n
    wr.interpret = interp({"name": "Anna Rossi"}, "Rifacimento bagno")
    r = await server._generate_work_report(me, r["text"], client_item_id="item_r2")
    assert r["client_item_id"] == "item_r2" and r["commessa_created"]
    bag = next(s for s in db.collection_sub_items.docs if s["item_id"] == "item_r2")
    assert bag["data"]["indirizzo_cantiere"] == "Corso Francia 1"   # client's address when no site was said
    wr.interpret = interp({"name": "Luigi Rossi"}, "Tettoia")
    r = await server._generate_work_report(me, "sopralluogo da Rossi per la tettoia", new_client=True)
    assert r["client_created"] and r["client_name"] == "Luigi Rossi"

    # 4) no client said: the report is still made, saved under "Report"
    wr.interpret = interp({}, "")
    nitems = len(db.collection_items.docs)
    r = await server._generate_work_report(me, "verifica quadro elettrico condominiale")
    assert r["client_item_id"] is None and SAVED[-1]["folder"] == "Report" and SAVED[-1]["subs"] == []
    assert len(db.collection_items.docs) == nitems and "senza cliente" in r["docx_filename"]

    # 5) the archive and the download
    lst = await server.list_work_reports(me)
    assert len(lst) == 5   # (the fake DB ignores projections: docx_b64 exclusion is Mongo's job)
    resp = await server.download_work_report(lst[0]["id"], me)
    assert resp.body[:2] == b"PK" and "filename*=UTF-8''" in resp.headers["content-disposition"]
    t = await server.list_work_templates(me)
    assert t["builtin"] == [{"key": "sopralluogo", "name": "Primo sopralluogo"}]
    # 6) folder names Drive/OneDrive accept
    assert server._drive_safe('Bagno 2/3: "nuovo"') == "Bagno 2 3 nuovo"
    # 7) Telegram: /report for an artigiano -> work templates; ambiguity -> buttons with "cliente nuovo"
    import telegram_bot as tb
    class Msg:
        def __init__(s, chat): s.chat = type("C", (), {"id": chat})(); s.sent = []
        async def reply_text(s, text, **k): SENT.append((text, k.get("reply_markup")))
    class Bot:
        async def send_chat_action(s, **k): pass
        async def send_message(s, chat_id, text, reply_markup=None, **k): SENT.append((text, reply_markup))
    class Ctx: bot = Bot()
    SENT = []
    await db.users.update_one({"user_id": "u1"}, {"$set": {"telegram_chat_id": 77, "business_vertical": "artigiano",
                                                          "org_id": me.org_id, "org_role": "owner"}})
    upd = type("U", (), {})(); upd.message = Msg(77); upd.effective_chat = upd.message.chat
    await tb._cmd_report(upd, Ctx())
    kb = SENT[-1][1].inline_keyboard
    assert kb[0][0].callback_data == "wrkrep:sopralluogo", kb
    user = await db.users.find_one({"user_id": "u1"})
    await tb._set_state(db, 77, "u1", pending_report_type="work:sopralluogo")
    st = await tb._get_state(db, "u1", 77)
    await tb._run_report_flow(upd, Ctx(), db, user, "sopralluogo da Rossi per il cancello", st["pending_report_type"])
    print(SENT[-1][0]); assert "più clienti" in SENT[-1][0]
    datas = [r[0].callback_data for r in SENT[-1][1].inline_keyboard]
    assert "wrkcli:new" in datas and "wrkcli:item_r2" in datas
    st = await tb._get_state(db, "u1", 77)
    assert st["pending_report_context"]["work"] and not st.get("pending_report_type")
    wr.interpret = interp({"name": "Anna Rossi"}, "Cancello automatico")
    await tb._run_work_report_flow(upd, Ctx(), db, user, st["pending_report_context"]["text"], "sopralluogo", client_item_id="item_r2")
    print(SENT[-1][0]); assert "Anna Rossi" in SENT[-1][0] and "Cancello automatico" in SENT[-1][0]
    # a vet (or no vertical) still gets the vet templates
    await db.users.update_one({"user_id": "u1"}, {"$set": {"business_vertical": "veterinario"}})
    await tb._cmd_report(upd, Ctx())
    assert SENT[-1][1].inline_keyboard[0][0].callback_data.startswith("vetrep:")
    print("ALL OK")
asyncio.run(main())
