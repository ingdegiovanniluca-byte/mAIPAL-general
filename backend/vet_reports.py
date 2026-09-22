"""Report generation for the "veterinario" vertical: turns a dictated/written visit
note into a structured .docx report, following a built-in or a custom uploaded
template, with patient data auto-filled from the "Pazienti" list when recognized.

Design choice: generated reports use one consistent, clean layout (heading + a
field/value table per section) rather than trying to byte-for-byte clone an
uploaded template's original formatting (checkboxes, merged cells, etc.) - the
uploaded template is used as a *structural* reference (which sections, which
fields), not a pixel-perfect stencil.
"""
import io
import json
import logging
import os
import re
from datetime import datetime, timezone
from typing import Optional

import openai
from docx import Document
from docx.shared import Pt

import usage_tracking as ut

logger = logging.getLogger(__name__)

MODEL = "gpt-4o"
FEATURE = "creazione_report"


def _track(resp, user_id: Optional[str], channel: str, status: str = "ok"):
    """This module builds its own openai client (bypasses LlmChat), so it calls
    usage_tracking directly right after chat.completions.create(...) - same
    fire-and-forget, never-block, never-raise pattern LlmChat uses internally."""
    usage = getattr(resp, "usage", None) if resp is not None else None
    cached = 0
    details = getattr(usage, "prompt_tokens_details", None) if usage else None
    if details is not None:
        cached = getattr(details, "cached_tokens", 0) or 0
    ut.fire_and_forget_llm_call(
        user_id=user_id, feature=FEATURE, channel=channel, trigger="utente",
        model=(getattr(resp, "model", None) if resp is not None else None) or MODEL,
        endpoint="chat.completions",
        input_tokens=(getattr(usage, "prompt_tokens", 0) or 0) if usage else 0,
        cached_input_tokens=cached,
        output_tokens=(getattr(usage, "completion_tokens", 0) or 0) if usage else 0,
        status=status, request_id=getattr(resp, "id", None) if resp is not None else None,
    )

# ---- Built-in templates ----------------------------------------------------
IMAGING_SECTIONS = [
    {"title": "Dati identificativi - proprietario e paziente", "fields": [
        "Cognome/Nome proprietario", "Telefono", "Email",
        "Nome animale", "Specie", "Razza", "Sesso", "Data di nascita", "Età",
        "Peso (kg)", "Microchip n°",
    ]},
    {"title": "Dati dell'esame", "fields": ["Data esame", "Veterinario", "Motivo esame"]},
    {"title": "Reperto ecografico", "fields": [
        "Fegato", "Milza", "Reni", "Vescica", "Surreni", "App. gastroenterico",
        "Pancreas", "Peritoneo", "Linfonodi", "Testicoli", "Prostata", "Ovaia", "Utero",
    ]},
    {"title": "Diagnosi ecografica e conclusioni", "fields": ["Conclusioni", "Referto complessivo"]},
    {"title": "Raccomandazioni e follow-up", "fields": ["Approfondimenti diagnostici", "Controllo ecografico", "Note aggiuntive"]},
]

# Nessun template "visita generale" è stato fornito: questa è una struttura di default
# ragionevole per una visita veterinaria generica, da rivedere/adattare.
GENERAL_SECTIONS = [
    {"title": "Dati identificativi - proprietario e paziente", "fields": [
        "Cognome/Nome proprietario", "Telefono", "Email",
        "Nome animale", "Specie", "Razza", "Sesso", "Data di nascita", "Età",
        "Peso (kg)", "Microchip n°",
    ]},
    {"title": "Dati della visita", "fields": ["Data visita", "Veterinario", "Motivo della visita"]},
    {"title": "Anamnesi", "fields": ["Anamnesi remota", "Anamnesi prossima"]},
    {"title": "Esame obiettivo generale", "fields": [
        "Stato generale", "Temperatura", "Frequenza cardiaca", "Frequenza respiratoria",
        "Mucose", "Linfonodi", "Cute e mantello", "Apparato cardiovascolare",
        "Apparato respiratorio", "Apparato gastroenterico", "Apparato muscoloscheletrico",
    ]},
    {"title": "Diagnosi e valutazione", "fields": ["Diagnosi presuntiva/definitiva", "Esami consigliati"]},
    {"title": "Terapia e raccomandazioni", "fields": ["Terapia prescritta", "Note aggiuntive", "Prossimo controllo"]},
]

BUILTIN_TEMPLATES = {
    "imaging": {"name": "Diagnostica per immagini", "sections": IMAGING_SECTIONS},
    "general": {"name": "Visita generale", "sections": GENERAL_SECTIONS},
}

# Mappa: chiave del campo paziente (Pazienti list) -> etichetta del campo nella
# sezione "Dati identificativi" del report. Usata per pre-compilare i dati
# anagrafici in modo affidabile (senza fidarsi ciecamente dell'LLM su questi campi).
PATIENT_FIELD_TO_REPORT_LABEL = {
    "proprietario": "Cognome/Nome proprietario",
    "telefono_proprietario": "Telefono",
    "email_proprietario": "Email",
    "nome": "Nome animale",
    "tipo_animale": "Specie",
    "razza": "Razza",
    "sesso": "Sesso",
    "data_nascita": "Data di nascita",
    "peso": "Peso (kg)",
    "microchip": "Microchip n°",
}

PATIENT_LIST_NAME = "Pazienti"
PATIENT_FIELDS = [
    {"key": "nome", "label": "Nome del paziente", "type": "text"},
    {"key": "tipo_animale", "label": "Tipo di animale", "type": "text"},
    {"key": "razza", "label": "Razza", "type": "text"},
    {"key": "sesso", "label": "Sesso", "type": "select", "options": ["maschio", "femmina"]},
    {"key": "stato_riproduttivo", "label": "Stato riproduttivo", "type": "select", "options": ["sterilizzato", "non sterilizzato"]},
    {"key": "microchip", "label": "Microchip", "type": "text"},
    {"key": "data_nascita", "label": "Data di nascita", "type": "date"},
    {"key": "peso", "label": "Peso", "type": "text"},
    {"key": "proprietario", "label": "Proprietario (nome)", "type": "text"},
    {"key": "email_proprietario", "label": "Email del proprietario", "type": "email"},
    {"key": "telefono_proprietario", "label": "Numero di telefono del proprietario", "type": "phone"},
    {"key": "note_anamnestiche", "label": "Note anamnestiche / allergie note", "type": "textarea"},
    {"key": "data_ultima_visita", "label": "Data dell'ultima visita", "type": "date"},
]


# ---- Custom template upload: structural extraction -------------------------
def _iter_block_items(doc):
    """Yields each Paragraph/Table in a docx in document order (python-docx only
    exposes .paragraphs and .tables as separate, unordered-relative-to-each-other
    lists - this walks the underlying XML body to interleave them correctly)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


_SECTION_TITLE_RE = re.compile(r"^\s*\d+[\.\)]\s*(.+)$")


def extract_template_structure(path: str) -> list[dict]:
    """Best-effort structural read of an uploaded .docx: numbered/heading paragraphs
    become section titles, and each table's left-column cell text (or first cell of
    each row, for 2-col tables) becomes a field name under the nearest preceding
    section. Falls back to a single "Contenuto" section if nothing is recognized."""
    doc = Document(path)
    sections: list[dict] = []
    current = None

    def ensure_current():
        nonlocal current
        if current is None:
            current = {"title": "Contenuto", "fields": []}
            sections.append(current)
        return current

    for block in _iter_block_items(doc):
        cls_name = block.__class__.__name__
        if cls_name == "Paragraph":
            text = (block.text or "").strip()
            if not text:
                continue
            is_heading = bool(block.style and "Heading" in (block.style.name or ""))
            m = _SECTION_TITLE_RE.match(text)
            if is_heading or (m and len(text) < 120):
                title = m.group(1).strip() if m else text
                current = {"title": title, "fields": []}
                sections.append(current)
        elif cls_name == "Table":
            sec = ensure_current()
            rows = list(block.rows)
            # A multi-row table whose first row's cells carry no ":" is usually a
            # column header ("ORGANO | REPERTO"), not a label:value pair - skip it.
            if len(rows) > 2 and rows[0].cells and not (rows[0].cells[0].text or "").strip().endswith(":"):
                rows = rows[1:]
            for row in rows:
                cells = row.cells
                if not cells:
                    continue
                # A row can hold more than one label:value pair side by side
                # (e.g. "Specie: | cane | Sesso: | maschio") - walk it two cells
                # at a time instead of only ever reading the first column.
                for i in range(0, len(cells) - 1, 2):
                    raw = (cells[i].text or "").strip()
                    if not raw:
                        continue
                    label = raw.splitlines()[0].strip().rstrip(":")
                    if label and len(label) < 80 and label not in sec["fields"]:
                        sec["fields"].append(label)

    return [s for s in sections if s["fields"]] or [{"title": "Contenuto", "fields": ["Note"]}]


# ---- Report generation -------------------------------------------------------
def find_matching_patients(text: str, patients: list[dict]) -> list[dict]:
    """Deterministic (non-LLM) name matching: whole-word, case-insensitive search of
    each patient's registered name in the dictation. Several *distinct* patients
    matching (e.g. two dogs both named "Fester") signals real ambiguity the caller
    must resolve before generating anything - this is why matching happens as an
    explicit pre-step instead of being left to the model to silently pick one."""
    text_low = (text or "").lower()
    matches = []
    seen_ids = set()
    for p in patients:
        name = (p.get("data", {}).get("nome") or "").strip()
        if not name or p["id"] in seen_ids:
            continue
        if re.search(rf"\b{re.escape(name.lower())}\b", text_low):
            matches.append(p)
            seen_ids.add(p["id"])
    return matches


# Organ fields that only make anatomical sense for one sex - matched case-insensitively
# against template field labels, so this also applies to custom uploaded templates.
_FEMALE_ONLY_ORGANS = {"ovaia", "utero"}
_MALE_ONLY_ORGANS = {"testicoli", "prostata"}


def filter_sections_for_patient(sections_skeleton: list[dict], patient: dict | None) -> list[dict]:
    """Drops organ fields not relevant to this patient's sex/reproductive status (e.g.
    no 'Ovaia'/'Utero' for a male; no 'Testicoli'/'Prostata' for a female; sterilized/
    castrated patients also lose the reproductive organs removed by that surgery).
    Returns a new list - never mutates the shared BUILTIN_TEMPLATES constants."""
    if not patient:
        return [{"title": s["title"], "fields": list(s["fields"])} for s in sections_skeleton]

    data = patient.get("data", {})
    sesso = (data.get("sesso") or "").strip().lower()
    stato = (data.get("stato_riproduttivo") or "").strip().lower()
    sterilized = "steriliz" in stato or "castrat" in stato

    exclude = set()
    if sesso == "maschio":
        exclude |= _FEMALE_ONLY_ORGANS
        if sterilized:
            exclude |= {"testicoli"}
    elif sesso == "femmina":
        exclude |= _MALE_ONLY_ORGANS
        if sterilized:
            exclude |= _FEMALE_ONLY_ORGANS

    if not exclude:
        return [{"title": s["title"], "fields": list(s["fields"])} for s in sections_skeleton]
    return [
        {"title": s["title"], "fields": [f for f in s["fields"] if f.strip().lower() not in exclude]}
        for s in sections_skeleton
    ]


async def interpret_visit(text: str, sections_skeleton: list[dict], user_id: Optional[str] = None, channel: str = "web") -> dict:
    """Asks the LLM to structure the dictation according to the given (already
    sex-filtered) section/field skeleton, in technical veterinary register. Patient
    identification is NOT the model's job - it's resolved deterministically before
    this is called. Returns {"sections": [{"title":..., "fields": {label: value}}]}."""
    skeleton_desc = json.dumps([{"title": s["title"], "fields": s["fields"]} for s in sections_skeleton], ensure_ascii=False)
    system = (
        "Sei l'assistente di un medico veterinario. Ricevi il resoconto (dettato o scritto) di una visita "
        "e devi strutturarlo secondo uno schema di sezioni/campi dato, scrivendo come un referto veterinario "
        "professionale.\n\n"
        f"Schema sezioni/campi da compilare (rispetta ESATTAMENTE titoli e nomi campo):\n{skeleton_desc}\n\n"
        "Regole:\n"
        "1. Compila SOLO i campi per cui il testo fornisce informazioni reali: se qualcosa non è menzionato, lascialo "
        "come stringa vuota. NON inventare dati clinici, diagnosi o misurazioni non dette.\n"
        "2. Registro linguistico: terminologia clinica veterinaria tecnica e professionale (es. \"non si "
        "apprezzano alterazioni ecostrutturali di rilievo\" invece di \"tutto normale\", \"tumefazione\" invece di "
        "\"gonfiore\", \"algia\" invece di \"dolore\", nomenclatura anatomica corretta), come in un referto scritto "
        "da un veterinario per un collega, MAI un riassunto in linguaggio comune.\n"
        "3. Frasi impersonali e concise, niente prima persona, niente commenti fuori dai campi dello schema.\n"
        "Rispondi SOLO con un JSON valido, in questo formato esatto: "
        '{"sections": [{"title": "...", "fields": {"Nome campo": "valore", ...}}, ...]}'
    )
    client = openai.AsyncOpenAI(api_key=os.environ["OPENAI_API_KEY"])
    try:
        resp = await client.chat.completions.create(
            model=MODEL,
            max_completion_tokens=4096,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": text}],
        )
    except Exception:
        _track(None, user_id, channel, status="errore")
        raise
    _track(resp, user_id, channel)
    raw = resp.choices[0].message.content or ""
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if not match:
        raise ValueError("Nessun JSON nella risposta del modello")
    parsed = json.loads(match.group(0))

    by_title = {s["title"]: s for s in parsed.get("sections", []) if isinstance(s, dict)}
    normalized = []
    for skel in sections_skeleton:
        got = by_title.get(skel["title"], {})
        fields_out = {}
        got_fields = got.get("fields", {}) if isinstance(got, dict) else {}
        for f in skel["fields"]:
            fields_out[f] = str(got_fields.get(f) or "").strip()
        normalized.append({"title": skel["title"], "fields": fields_out})

    return {"sections": normalized}


def apply_patient_data(sections: list[dict], patient: dict | None) -> None:
    """Overwrites the identificativi fields in-place with authoritative patient data,
    when a patient was matched - clinical/free-text fields from the LLM are untouched."""
    if not patient:
        return
    data = patient.get("data", {})
    for sec in sections:
        for patient_key, label in PATIENT_FIELD_TO_REPORT_LABEL.items():
            if label in sec["fields"] and data.get(patient_key):
                sec["fields"][label] = str(data[patient_key])


def build_docx(title: str, patient_name: str | None, sections: list[dict]) -> bytes:
    doc = Document()
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.size = Pt(20)

    sub = doc.add_paragraph(f"Paziente: {patient_name}" if patient_name else "Paziente: non specificato")
    sub.runs[0].italic = True

    doc.add_paragraph(f"Data generazione: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC")

    for sec in sections:
        doc.add_heading(sec["title"], level=1)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for label, value in sec["fields"].items():
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value or "-"
            for p in row.cells[0].paragraphs:
                for r in p.runs:
                    r.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
