"""Reports for the "artigiano" vertical (primo sopralluogo, later rapporto d'intervento):
a dictated/written description becomes a structured .docx, the client and the commessa
are recognized in (or added to) the Clienti list.

Same approach as vet_reports.py: the model only fills a fixed skeleton of sections /
fields from what was said (never inventing measures, prices or dates); who the client is
is resolved deterministically by the caller, not left to the model.
"""
import io
import json
import logging
import os
import re
from datetime import datetime
from typing import Optional
from zoneinfo import ZoneInfo

import openai
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import usage_tracking as ut

logger = logging.getLogger(__name__)

MODEL = "gpt-4o"
FEATURE = "creazione_report"
LOCAL_TZ = ZoneInfo("Europe/Rome")

SOPRALLUOGO_SECTIONS = [
    {"title": "Dati del cliente", "fields": ["Cliente", "Telefono", "Email", "Indirizzo del cantiere"]},
    {"title": "Sopralluogo", "fields": ["Data", "Eseguito da", "Oggetto dell'intervento"]},
    {"title": "Stato attuale e rilievi", "fields": ["Descrizione dello stato attuale", "Misure e rilievi",
                                                  "Impianti / strutture esistenti", "Criticità riscontrate"]},
    {"title": "Proposta", "fields": ["Lavori proposti", "Materiali stimati", "Tempi indicativi", "Stima economica indicativa"]},
    {"title": "Note", "fields": ["Note e prescrizioni", "Prossimi passi"]},
]

BUILTIN_TEMPLATES = {
    "sopralluogo": {"name": "Primo sopralluogo", "sections": SOPRALLUOGO_SECTIONS},
}

# report label <- Clienti list base field
CLIENT_FIELD_TO_LABEL = {"nome": "Cliente", "telefono": "Telefono", "email": "Email"}


def _track(resp, user_id: Optional[str], channel: str, status: str = "ok"):
    usage = getattr(resp, "usage", None) if resp is not None else None
    ut.fire_and_forget_llm_call(
        user_id=user_id, feature=FEATURE, channel=channel, trigger="utente",
        model=(getattr(resp, "model", None) if resp is not None else None) or MODEL, endpoint="chat.completions",
        input_tokens=(getattr(usage, "prompt_tokens", 0) or 0) if usage else 0, cached_input_tokens=0,
        output_tokens=(getattr(usage, "completion_tokens", 0) or 0) if usage else 0,
        status=status, request_id=getattr(resp, "id", None) if resp is not None else None,
    )


def find_matching_clients(text: str, clients: list, name_key: str) -> list:
    """Clients whose name appears in the text as whole words ("Rossi" also finds "Mario
    Rossi" when it's the only Rossi). Several distinct ones = the caller asks."""
    low = (text or "").lower()
    full, partial = [], []
    for c in clients:
        name = str((c.get("data") or {}).get(name_key) or "").strip()
        if not name:
            continue
        if re.search(rf"\b{re.escape(name.lower())}\b", low):
            full.append(c)
            continue
        words = [w for w in re.findall(r"[\wàèéìòù']+", name.lower()) if len(w) >= 4]
        if words and any(re.search(rf"\b{re.escape(w)}\b", low) for w in words):
            partial.append(c)
    return full or partial


async def interpret(text: str, sections: list, user_name: str = "", user_id: Optional[str] = None, channel: str = "web") -> dict:
    """{"sections": [{"title", "fields": {label: value}}], "client": {...}, "job_title": "..."}"""
    today = datetime.now(LOCAL_TZ)
    skeleton = json.dumps([{"title": s["title"], "fields": s["fields"]} for s in sections], ensure_ascii=False)
    system = (
        "Sei l'assistente di un'impresa artigiana (impianti elettrici, idraulici, edilizia, manutenzioni). Ricevi "
        "il resoconto dettato o scritto di un intervento o di un sopralluogo e lo strutturi in un report "
        f"professionale da consegnare al cliente. Oggi è {today.strftime('%d/%m/%Y')}; il report lo firma {user_name or 'il titolare'}.\n\n"
        f"Schema da compilare (rispetta ESATTAMENTE titoli e nomi dei campi):\n{skeleton}\n\n"
        "Regole:\n"
        "1. Compila SOLO con ciò che è detto nel testo. Se un'informazione manca lascia la stringa vuota: NON inventare "
        "misure, quantità, prezzi, tempi o date. Se non c'è una data, usa la data di oggi solo nel campo 'Data'.\n"
        "2. Linguaggio tecnico e professionale del settore, frasi impersonali e concise (es. 'Si rileva…', "
        "'Si propone…'), niente prima persona, niente commenti fuori dai campi.\n"
        "3. Liste di lavori o materiali: una voce per riga, preceduta da '- '.\n"
        "Rispondi SOLO con JSON: {\"sections\": [{\"title\": \"...\", \"fields\": {\"Nome campo\": \"valore\"}}], "
        "\"client\": {\"name\": \"nome e cognome o ragione sociale del cliente, se detto\", \"phone\": \"\", \"email\": \"\", "
        "\"address\": \"indirizzo del cantiere, se detto\"}, \"job_title\": \"titolo breve dei lavori (max 8 parole), "
        "es. 'Rifacimento impianto elettrico cucina'\"}"
    )
    client = openai.AsyncOpenAI(api_key=os.environ["OPENAI_API_KEY"])
    try:
        resp = await client.chat.completions.create(
            model=MODEL, max_completion_tokens=4096, response_format={"type": "json_object"},
            messages=[{"role": "system", "content": system}, {"role": "user", "content": text}],
        )
    except Exception:
        _track(None, user_id, channel, status="errore")
        raise
    _track(resp, user_id, channel)
    m = re.search(r"\{.*\}", resp.choices[0].message.content or "", re.S)
    if not m:
        raise ValueError("Nessun JSON nella risposta del modello")
    parsed = json.loads(m.group(0))
    by_title = {s.get("title"): s for s in parsed.get("sections", []) if isinstance(s, dict)}
    out = []
    for sk in sections:
        got = (by_title.get(sk["title"]) or {}).get("fields") or {}
        out.append({"title": sk["title"], "fields": {f: str(got.get(f) or "").strip() for f in sk["fields"]}})
    cl = parsed.get("client") if isinstance(parsed.get("client"), dict) else {}
    return {
        "sections": out,
        "client": {k: str(cl.get(k) or "").strip() for k in ("name", "phone", "email", "address")},
        "job_title": str(parsed.get("job_title") or "").strip()[:80],
    }


def set_field(sections: list, label: str, value: str, overwrite: bool = True) -> None:
    for s in sections:
        if label in s["fields"] and value and (overwrite or not s["fields"][label]):
            s["fields"][label] = value


def build_docx(title: str, client_name: Optional[str], job_title: Optional[str], sections: list, company: str = "") -> bytes:
    doc = Document()
    if company:
        p = doc.add_paragraph(company)
        p.runs[0].bold = True
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.size = Pt(20)
    sub = doc.add_paragraph(" · ".join(x for x in (client_name and f"Cliente: {client_name}", job_title and f"Lavori: {job_title}") if x)
                            or "Cliente: non indicato")
    sub.runs[0].italic = True
    for sec in sections:
        if not any(v for v in sec["fields"].values()):
            continue   # a section nobody talked about stays out of the client's copy
        doc.add_heading(sec["title"], level=1)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for label, value in sec["fields"].items():
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value or "-"
            for p in row.cells[0].paragraphs:
                for r in p.runs:
                    r.bold = True
    # signatures: room to sign on the printed copy (no digital signature for now)
    doc.add_paragraph()
    doc.add_paragraph("Luogo e data: ________________________________")
    doc.add_paragraph()
    sig = doc.add_table(rows=2, cols=2)
    sig.cell(0, 0).text = "Firma del tecnico"
    sig.cell(0, 1).text = "Firma del cliente per presa visione"
    sig.cell(1, 0).text = "\n\n______________________________"
    sig.cell(1, 1).text = "\n\n______________________________"
    for c in (sig.cell(0, 0), sig.cell(0, 1)):
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
